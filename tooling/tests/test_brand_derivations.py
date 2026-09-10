"""
Tests for core/brand/derivations/ — color math, loader, PBI theme, CSS variables.

Uses aurora_group/brand/brand_spec.yaml as the integration fixture.
All converter tests run against that real spec so regressions on actual
brand output are caught without mocks.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

# Ensure repo root is on sys.path for imports.
_REPO_ROOT = Path(__file__).resolve().parents[2]
if str(_REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(_REPO_ROOT))

_AURORA_SPEC = _REPO_ROOT / "showcases" / "aurora_group" / "brand" / "brand_spec.yaml"

# ─────────────────────────────────────────────
# _color_math
# ─────────────────────────────────────────────

from core.brand.derivations._color_math import (
    parse_hex,
    to_hex,
    lighten,
    darken,
    rotate_hue,
    with_opacity_hex,
    with_opacity_rgba,
    derive_data_palette,
)


class TestParseHex:
    def test_six_digit(self):
        assert parse_hex("#2ECDE7") == (0x2E, 0xCD, 0xE7)

    def test_three_digit_expands(self):
        assert parse_hex("#FFF") == (255, 255, 255)

    def test_case_insensitive(self):
        assert parse_hex("#2ecde7") == parse_hex("#2ECDE7")

    def test_invalid_raises(self):
        with pytest.raises(ValueError):
            parse_hex("#GGGGGG")

    def test_no_hash(self):
        assert parse_hex("2ECDE7") == (0x2E, 0xCD, 0xE7)


class TestToHex:
    def test_round_trip(self):
        r, g, b = parse_hex("#2ECDE7")
        assert to_hex(r, g, b) == "#2ECDE7"

    def test_clamps_max(self):
        result = to_hex(300, 300, 300)
        assert result == "#FFFFFF"

    def test_clamps_min(self):
        result = to_hex(-10, -10, -10)
        assert result == "#000000"


class TestLighten:
    def test_returns_hex(self):
        result = lighten("#107C10", 0.3)
        assert result.startswith("#")
        assert len(result) == 7

    def test_lightened_is_different(self):
        assert lighten("#107C10", 0.3) != "#107C10"

    def test_full_lighten_approaches_white(self):
        result = lighten("#107C10", 1.0)
        r, g, b = parse_hex(result)
        assert r > 240 and g > 240 and b > 240

    def test_zero_amount_unchanged(self):
        assert lighten("#107C10", 0.0) == "#107C10"


class TestDarken:
    def test_darkened_is_different(self):
        assert darken("#2ECDE7", 0.3) != "#2ECDE7"

    def test_full_darken_approaches_black(self):
        result = darken("#2ECDE7", 1.0)
        r, g, b = parse_hex(result)
        assert r < 15 and g < 15 and b < 15

    def test_zero_amount_unchanged(self):
        assert darken("#2ECDE7", 0.0) == "#2ECDE7"


class TestRotateHue:
    def test_360_returns_same(self):
        original = "#2ECDE7"
        assert rotate_hue(original, 360) == original

    def test_180_is_different(self):
        assert rotate_hue("#2ECDE7", 180) != "#2ECDE7"

    def test_returns_hex_format(self):
        result = rotate_hue("#2ECDE7", 30)
        assert result.startswith("#") and len(result) == 7


class TestWithOpacity:
    def test_hex_format(self):
        result = with_opacity_hex("#2ECDE7", 0.10)
        assert result.startswith("#") and len(result) == 9

    def test_rgba_format(self):
        result = with_opacity_rgba("#2ECDE7", 0.10)
        assert result.startswith("rgba(")
        assert "0.10" in result

    def test_full_opacity_hex(self):
        result = with_opacity_hex("#FF0000", 1.0)
        assert result.endswith("FF")

    def test_zero_opacity_hex(self):
        result = with_opacity_hex("#FF0000", 0.0)
        assert result.endswith("00")


class TestDeriveDataPalette:
    def test_returns_eight_colors(self):
        palette = derive_data_palette("#2ECDE7", "#44B396")
        assert len(palette) == 8

    def test_all_valid_hex(self):
        palette = derive_data_palette("#2ECDE7", "#44B396")
        for color in palette:
            assert color.startswith("#"), f"Not hex: {color}"
            assert len(color) == 7, f"Wrong length: {color}"

    def test_primary_is_first(self):
        palette = derive_data_palette("#2ECDE7", "#44B396", "NeutralAccent")
        assert palette[0] == "#2ECDE7"

    def test_secondary_is_second_for_neutral_accent(self):
        palette = derive_data_palette("#2ECDE7", "#44B396", "NeutralAccent")
        assert palette[1] == "#44B396"

    def test_monochromatic_all_same_hue(self):
        import colorsys
        from core.brand.derivations._color_math import parse_hex
        palette = derive_data_palette("#2ECDE7", "#44B396", "Monochromatic")
        primary_h = colorsys.rgb_to_hls(*[x / 255 for x in parse_hex("#2ECDE7")])[0]
        for color in palette:
            h = colorsys.rgb_to_hls(*[x / 255 for x in parse_hex(color)])[0]
            assert abs(h - primary_h) < 0.02, f"Hue shifted for {color}: {h} vs {primary_h}"

    def test_divergent_returns_eight(self):
        palette = derive_data_palette("#2ECDE7", "#44B396", "Divergent")
        assert len(palette) == 8

    def test_unknown_concept_defaults_to_neutral_accent(self):
        palette = derive_data_palette("#2ECDE7", "#44B396", "Unknown")
        assert len(palette) == 8
        assert palette[0] == "#2ECDE7"


# ─────────────────────────────────────────────
# loader
# ─────────────────────────────────────────────

from core.brand.derivations.loader import load_brand_spec


class TestLoader:
    def test_loads_aurora_spec(self):
        spec = load_brand_spec(_AURORA_SPEC)
        assert isinstance(spec, dict)

    def test_identity_present(self):
        spec = load_brand_spec(_AURORA_SPEC)
        assert spec["identity"]["brand_id"] == "aurora_group"

    def test_primary_color_present(self):
        spec = load_brand_spec(_AURORA_SPEC)
        assert spec["color"]["primary"] == "#2ECDE7"

    def test_secondary_color_present(self):
        spec = load_brand_spec(_AURORA_SPEC)
        assert spec["color"]["secondary"] == "#44B396"

    def test_neutral_scale_has_six_steps(self):
        spec = load_brand_spec(_AURORA_SPEC)
        ns = spec["color"]["neutral_scale"]
        assert set(ns.keys()) >= {"50", "100", "200", "400", "700", "900"}

    def test_file_not_found(self):
        with pytest.raises(FileNotFoundError):
            load_brand_spec(Path("/nonexistent/brand_spec.yaml"))

    def test_missing_field_raises_value_error(self, tmp_path):
        bad_spec = tmp_path / "bad.yaml"
        bad_spec.write_text("identity:\n  brand_id: test\n", encoding="utf-8")
        with pytest.raises(ValueError, match="missing required fields"):
            load_brand_spec(bad_spec)

    def test_non_mapping_raises_value_error(self, tmp_path):
        bad_spec = tmp_path / "bad.yaml"
        bad_spec.write_text("- item1\n- item2\n", encoding="utf-8")
        with pytest.raises(ValueError):
            load_brand_spec(bad_spec)


# ─────────────────────────────────────────────
# pbi_theme
# ─────────────────────────────────────────────

from core.brand.derivations.pbi_theme import spec_to_pbi_theme


@pytest.fixture(scope="module")
def aurora_spec():
    return load_brand_spec(_AURORA_SPEC)


@pytest.fixture(scope="module")
def aurora_pbi_theme(aurora_spec):
    return spec_to_pbi_theme(aurora_spec)


class TestPbiTheme:
    def test_returns_dict(self, aurora_pbi_theme):
        assert isinstance(aurora_pbi_theme, dict)

    def test_has_name(self, aurora_pbi_theme):
        assert "name" in aurora_pbi_theme
        assert "Aurora Group SE" in aurora_pbi_theme["name"]

    def test_has_schema(self, aurora_pbi_theme):
        assert "$schema" in aurora_pbi_theme

    def test_data_colors_count(self, aurora_pbi_theme):
        assert len(aurora_pbi_theme["dataColors"]) == 8

    def test_data_colors_valid_hex(self, aurora_pbi_theme):
        for color in aurora_pbi_theme["dataColors"]:
            assert color.startswith("#"), f"Not hex: {color}"

    def test_primary_is_first_data_color(self, aurora_spec, aurora_pbi_theme):
        assert aurora_pbi_theme["dataColors"][0] == aurora_spec["color"]["primary"]

    def test_semantic_good_matches_positive(self, aurora_spec, aurora_pbi_theme):
        assert aurora_pbi_theme["good"] == aurora_spec["color"]["semantic"]["positive"]["color"]

    def test_semantic_bad_matches_negative(self, aurora_spec, aurora_pbi_theme):
        assert aurora_pbi_theme["bad"] == aurora_spec["color"]["semantic"]["negative"]["color"]

    def test_semantic_neutral_matches_neutral(self, aurora_spec, aurora_pbi_theme):
        assert aurora_pbi_theme["neutral"] == aurora_spec["color"]["semantic"]["neutral"]["color"]

    def test_background_is_neutral_100(self, aurora_spec, aurora_pbi_theme):
        assert aurora_pbi_theme["background"] == aurora_spec["color"]["neutral_scale"]["100"]

    def test_foreground_is_neutral_900(self, aurora_spec, aurora_pbi_theme):
        assert aurora_pbi_theme["foreground"] == aurora_spec["color"]["neutral_scale"]["900"]

    def test_has_visual_styles(self, aurora_pbi_theme):
        assert "visualStyles" in aurora_pbi_theme
        assert isinstance(aurora_pbi_theme["visualStyles"], dict)

    def test_maximum_is_lightened_positive(self, aurora_spec, aurora_pbi_theme):
        pos = aurora_spec["color"]["semantic"]["positive"]["color"]
        assert aurora_pbi_theme["maximum"] != pos  # lightened — must differ

    def test_production_canvas_adds_delta(self, aurora_spec):
        base = spec_to_pbi_theme(aurora_spec, canvas_profile="powerbi_design_base")
        prod = spec_to_pbi_theme(aurora_spec, canvas_profile="powerbi_production")
        # Production name should include the profile key
        assert "powerbi_production" in prod["name"]
        assert "powerbi_design_base" in base["name"]

    def test_neutral_accent_concept(self, aurora_spec):
        theme = spec_to_pbi_theme(aurora_spec, concept="NeutralAccent")
        assert theme["dataColors"][0] == aurora_spec["color"]["primary"]
        assert theme["dataColors"][1] == aurora_spec["color"]["secondary"]

    def test_is_serialisable_as_json(self, aurora_pbi_theme):
        dumped = json.dumps(aurora_pbi_theme)
        reloaded = json.loads(dumped)
        assert reloaded["dataColors"] == aurora_pbi_theme["dataColors"]


# ─────────────────────────────────────────────
# css_variables
# ─────────────────────────────────────────────

from core.brand.derivations.css_variables import spec_to_css


@pytest.fixture(scope="module")
def aurora_css(aurora_spec):
    return spec_to_css(aurora_spec)


class TestCssVariables:
    def test_returns_string(self, aurora_css):
        assert isinstance(aurora_css, str)

    def test_has_root_block(self, aurora_css):
        assert ":root {" in aurora_css
        assert "}" in aurora_css

    def test_primary_color_present(self, aurora_spec, aurora_css):
        assert aurora_spec["color"]["primary"] in aurora_css

    def test_secondary_color_present(self, aurora_spec, aurora_css):
        assert aurora_spec["color"]["secondary"] in aurora_css

    def test_neutral_scale_50_present(self, aurora_spec, aurora_css):
        assert aurora_spec["color"]["neutral_scale"]["50"] in aurora_css

    def test_neutral_scale_900_present(self, aurora_spec, aurora_css):
        assert aurora_spec["color"]["neutral_scale"]["900"] in aurora_css

    def test_semantic_colors_present(self, aurora_spec, aurora_css):
        sem = aurora_spec["color"]["semantic"]
        assert sem["positive"]["color"] in aurora_css
        assert sem["negative"]["color"] in aurora_css
        assert sem["warning"]["color"] in aurora_css
        assert sem["neutral"]["color"] in aurora_css

    def test_font_primary_present(self, aurora_spec, aurora_css):
        assert "--brand-font-primary" in aurora_css

    def test_type_scale_steps_present(self, aurora_css):
        for step in ("xs", "sm", "md", "lg", "xl", "2xl", "3xl"):
            assert f"--brand-type-{step}" in aurora_css, f"Missing --brand-type-{step}"

    def test_role_aliases_present(self, aurora_css):
        for role in ("kpi-value", "kpi-label", "body", "heading-1", "chart-title"):
            assert f"--brand-type-{role}" in aurora_css, f"Missing role --brand-type-{role}"

    def test_spacing_variables_present(self, aurora_css):
        assert "--brand-spacing-md" in aurora_css
        assert "--brand-spacing-xl" in aurora_css

    def test_border_variables_present(self, aurora_css):
        assert "--brand-radius-md" in aurora_css
        assert "--brand-border-medium" in aurora_css

    def test_shadow_variables_present(self, aurora_css):
        assert "--brand-shadow-low" in aurora_css
        assert "--brand-shadow-high" in aurora_css

    def test_semantic_bg_colors_present(self, aurora_css):
        assert "--brand-color-positive-bg" in aurora_css
        assert "--brand-color-negative-bg" in aurora_css

    def test_semantic_role_aliases_present(self, aurora_css):
        assert "--brand-color-background" in aurora_css
        assert "--brand-color-surface" in aurora_css
        assert "--brand-color-border" in aurora_css

    def test_clamp_syntax_used(self, aurora_css):
        assert "clamp(" in aurora_css

    def test_rgba_opacity_used_for_subtle(self, aurora_css):
        assert "rgba(" in aurora_css

    def test_source_hint_in_header(self, aurora_spec):
        css = spec_to_css(aurora_spec, source_hint="showcases/aurora_group/brand/brand_spec.yaml")
        assert "showcases/aurora_group/brand/brand_spec.yaml" in css

    def test_brand_name_in_header(self, aurora_spec, aurora_css):
        assert "Aurora Group SE" in aurora_css


# ─────────────────────────────────────────────
# docx_document (I-10.4)
# ─────────────────────────────────────────────

import io

from docx import Document
from docx.document import Document as DocxDocumentType
from docx.oxml.ns import qn

from core.brand.derivations.docx_document import DocSection, DocTable, HandoverDoc, spec_to_docx


@pytest.fixture(scope="module")
def sample_handover_doc():
    return HandoverDoc(
        title="Sample Report — Handover Documentation",
        subtitle_lines=["Test Product · a tagline", "Layer-Tool: report-documenter"],
        sections=[
            DocSection(heading="Business", level=1, paragraphs=["Report Pages:", "Overview — 2 visual(s): card, table"]),
            DocSection(heading="KPIs & Meaning", level=2,
                       table=DocTable(headers=["Measure", "Purpose", "Format"], rows=[["Net Sales", "Total sales", "#,0"]])),
            DocSection(heading="Row-Level Security", level=2, paragraphs=["(no roles)"]),
        ],
    )


@pytest.fixture(scope="module")
def aurora_docx(aurora_spec, sample_handover_doc):
    return spec_to_docx(aurora_spec, sample_handover_doc)


def _reopen(document) -> Document:
    """Round-trip a docx.Document through bytes, like a real consumer would."""
    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)
    return Document(buf)


class TestDocxDocument:
    def test_returns_document(self, aurora_docx):
        assert isinstance(aurora_docx, DocxDocumentType)

    def test_round_trips_through_save_and_reopen(self, aurora_docx):
        # OOXML bytes aren't stable run-to-run (timestamps, relationship IDs) —
        # the real determinism guarantee is that saved bytes reopen cleanly.
        reopened = _reopen(aurora_docx)
        assert len(reopened.paragraphs) > 0

    def test_title_present(self, aurora_docx):
        reopened = _reopen(aurora_docx)
        assert reopened.paragraphs[0].text == "Sample Report — Handover Documentation"

    def test_brand_name_present(self, aurora_spec, aurora_docx):
        reopened = _reopen(aurora_docx)
        texts = [p.text for p in reopened.paragraphs]
        assert aurora_spec["identity"]["brand_name"] in texts

    def test_section_headings_present(self, aurora_docx):
        reopened = _reopen(aurora_docx)
        headings = [p.text for p in reopened.paragraphs if p.style.name.startswith("Heading")]
        assert "Business" in headings
        assert "KPIs & Meaning" in headings
        assert "Row-Level Security" in headings

    def test_table_rendered_with_headers_and_rows(self, aurora_docx):
        reopened = _reopen(aurora_docx)
        assert len(reopened.tables) == 1
        table = reopened.tables[0]
        assert [c.text for c in table.rows[0].cells] == ["Measure", "Purpose", "Format"]
        assert [c.text for c in table.rows[1].cells] == ["Net Sales", "Total sales", "#,0"]

    def test_page_margins_set(self, aurora_docx):
        reopened = _reopen(aurora_docx)
        section = reopened.sections[0]
        assert section.left_margin == section.right_margin == section.top_margin == section.bottom_margin
        assert section.left_margin.inches == pytest.approx(1.0)

    def test_hyphenation_enabled(self, aurora_docx):
        reopened = _reopen(aurora_docx)
        el = reopened.settings.element.find(qn("w:autoHyphenation"))
        assert el is not None
        assert el.get(qn("w:val")) == "1"

    def test_widow_control_set_on_body_paragraphs(self, aurora_docx):
        reopened = _reopen(aurora_docx)
        body_paragraphs = [p for p in reopened.paragraphs if p.text and not p.style.name.startswith("Heading")]
        assert body_paragraphs, "expected at least one body paragraph"
        for p in body_paragraphs:
            assert p.paragraph_format.widow_control is True

    def test_title_uses_brand_primary_color(self, aurora_spec, aurora_docx):
        reopened = _reopen(aurora_docx)
        title_run = reopened.paragraphs[0].runs[0]
        assert str(title_run.font.color.rgb) == aurora_spec["color"]["primary"].lstrip("#")

    def test_body_font_matches_brand_typography(self, aurora_spec, aurora_docx):
        reopened = _reopen(aurora_docx)
        expected_font = aurora_spec["typography"]["font_family"]["primary"].split(",")[0].strip()
        assert reopened.styles["Normal"].font.name == expected_font

    def test_different_brand_produces_different_title_color(self, sample_handover_doc):
        generic_spec = load_brand_spec(_REPO_ROOT / "core" / "brand" / "samples" / "generic_brand.yaml")
        aurora_spec_local = load_brand_spec(_AURORA_SPEC)
        assert generic_spec["color"]["primary"] != aurora_spec_local["color"]["primary"]

        generic_doc = spec_to_docx(generic_spec, sample_handover_doc)
        aurora_doc = spec_to_docx(aurora_spec_local, sample_handover_doc)
        generic_color = str(_reopen(generic_doc).paragraphs[0].runs[0].font.color.rgb)
        aurora_color = str(_reopen(aurora_doc).paragraphs[0].runs[0].font.color.rgb)
        assert generic_color != aurora_color


# ─────────────────────────────────────────────
# Integration: derive_brand_artifacts (orchestrator)
# ─────────────────────────────────────────────

from tooling.brand.derive_brand_artifacts import derive_from_spec


class TestDeriveFromSpec:
    """
    Each test uses an isolated copy of the aurora spec so _patch_tool_derivations
    never writes /tmp paths back into the real brand_spec.yaml.
    """

    @pytest.fixture(autouse=True)
    def spec_copy(self, tmp_path):
        import shutil
        copy = tmp_path / "brand_spec.yaml"
        shutil.copy(_AURORA_SPEC, copy)
        self._spec = copy
        self._tmp = tmp_path

    def test_writes_pbi_theme_file(self):
        outputs = derive_from_spec(
            spec_path=self._spec,
            out_dir_pbi=self._tmp / "themes",
            out_dir_css=self._tmp / "css",
            concept="Monochromatic",
            verbose=False,
        )
        assert outputs["pbi_theme"].exists()
        assert outputs["pbi_theme"].suffix == ".json"

    def test_writes_css_file(self):
        outputs = derive_from_spec(
            spec_path=self._spec,
            out_dir_pbi=self._tmp / "themes",
            out_dir_css=self._tmp / "css",
            verbose=False,
        )
        assert outputs["css_variables"].exists()
        assert outputs["css_variables"].suffix == ".css"

    def test_pbi_json_is_valid(self):
        outputs = derive_from_spec(
            spec_path=self._spec,
            out_dir_pbi=self._tmp / "themes",
            out_dir_css=self._tmp / "css",
            verbose=False,
        )
        data = json.loads(outputs["pbi_theme"].read_text(encoding="utf-8"))
        assert "dataColors" in data
        assert len(data["dataColors"]) == 8

    def test_css_file_contains_root(self):
        outputs = derive_from_spec(
            spec_path=self._spec,
            out_dir_pbi=self._tmp / "themes",
            out_dir_css=self._tmp / "css",
            verbose=False,
        )
        css = outputs["css_variables"].read_text(encoding="utf-8")
        assert ":root {" in css

    def test_patches_tool_derivations(self):
        import yaml
        derive_from_spec(
            spec_path=self._spec,
            out_dir_pbi=self._tmp / "themes",
            out_dir_css=self._tmp,
            verbose=False,
        )
        patched = yaml.safe_load(self._spec.read_text(encoding="utf-8"))
        assert patched["tool_derivations"]["css_variables"] is not None
        assert patched["tool_derivations"]["powerbi_theme"] is not None

    def test_nonexistent_spec_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            derive_from_spec(
                spec_path=tmp_path / "missing.yaml",
                out_dir_pbi=tmp_path,
                out_dir_css=tmp_path,
                verbose=False,
            )
