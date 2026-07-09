"""Tests for the Report-Documenter Layer-Tool (task I-5.2).

DoD: a deterministic Doc-Gen snapshot (a checked-in golden Markdown that turns
red if the branded layout drifts) + standalone generation against the bare stack.
"""
from __future__ import annotations

from pathlib import Path

from tooling.superversion.canonical_contract import (
    CanonicalModel,
    Measure,
    ReportModel,
    ReportPage,
    Role,
    RoleTablePermission,
    SemanticModel,
    Table,
    Visual,
)
from tooling.superversion.from_aluca import from_bracket_file
from tooling.superversion.layer_tools import report_documenter as doc

REPO = Path(__file__).resolve().parents[3]
KPIS = REPO / "core/kpi_catalog/kpis"
COM001 = REPO / "core/usecases/core/COM-001_Sales_Performance/UseCase_Bracket.yaml"
GOLDEN = Path(__file__).resolve().parent / "golden_docs" / "COM-001.md"


def _com001():
    return from_bracket_file(COM001, KPIS)


def test_matches_golden_snapshot():
    """Regression snapshot: branded layout drift turns this red.
    Regenerate via: python -m tooling.superversion.layer_tools.report_documenter \\
        --out tooling/superversion/tests/golden_docs/COM-001.md"""
    md = doc.render_markdown(_com001())
    assert md == GOLDEN.read_text(encoding="utf-8"), "doc drifted from golden — regenerate if intended"


def test_deterministic():
    m = _com001()
    assert doc.render_markdown(m) == doc.render_markdown(m)


def test_has_business_and_technical_sections():
    md = doc.render_markdown(_com001())
    for marker in ("# COM-001", "## Business", "### KPIs & Meaning",
                   "## Technical", "### Measures (DAX dialect / HITL)", "### Row-Level Security"):
        assert marker in md, f"missing section: {marker}"


def test_branding_title_block():
    md = doc.render_markdown(_com001(), branding=doc.Branding(product_name="ACME BI", tagline="x"))
    assert "**ACME BI**" in md


def test_lists_every_page_and_measure():
    m = _com001()
    md = doc.render_markdown(m)
    for p in m.report.pages:
        assert (p.display_name or p.name) in md
    for t in m.semantic.tables:
        for measure in t.measures:
            assert measure.name in md


def test_rls_section_reflects_real_filter():
    m = CanonicalModel(
        semantic=SemanticModel(name="S", tables=[Table(name="f", measures=[Measure(name="M")])],
                               roles=[Role(name="org",
                                           table_permissions=[RoleTablePermission(
                                               table="f", filter_expression="[R]=1")])]),
        report=ReportModel(name="R", pages=[ReportPage(name="P", visuals=[Visual(visual_id="v", visual_type="card")])]),
    )
    md = doc.render_markdown(m)
    assert "**org**" in md and "[R]=1" in md


def test_hitl_dax_placeholder_shown():
    md = doc.render_markdown(_com001())
    assert "HITL: DAX dialect to be defined" in md


def test_cli_writes_file(tmp_path: Path):
    out = tmp_path / "doc.md"
    assert doc.main([str(COM001), "--out", str(out)]) == 0
    assert out.exists() and out.read_text(encoding="utf-8").startswith("# COM-001")


def test_cli_missing_bracket_returns_one():
    assert doc.main([str(REPO / "nope.yaml")]) == 1


# --------------------------------------------------------------------------- #
# DOCX (I-10.4)                                                               #
# --------------------------------------------------------------------------- #

import io

import pytest
from docx import Document


def _reopen(data: bytes) -> Document:
    return Document(io.BytesIO(data))


def test_md_and_docx_agree_on_content():
    """MD and DOCX are both formatters over the one shared _extract_content —
    this pins that they actually stay in sync (measures, tables, RLS) rather
    than silently drifting as two hand-maintained walks of CanonicalModel."""
    m = _com001()
    md = doc.render_markdown(m)
    docx_paragraphs = [p.text for p in _reopen(doc.render_docx(m)).paragraphs]
    docx_cells = [c.text for t in _reopen(doc.render_docx(m)).tables for r in t.rows for c in r.cells]

    for table in m.semantic.tables:
        for measure in table.measures:
            assert measure.name in md
            assert measure.name in docx_cells

    for p in m.report.pages:
        label = p.display_name or p.name
        assert label in md
        assert any(label in text for text in docx_paragraphs)

    for r in m.semantic.roles:
        if any((tp.filter_expression or "").strip() for tp in r.table_permissions):
            assert r.name in md
            assert any(r.name in text for text in docx_paragraphs)


def test_render_docx_returns_nonempty_bytes():
    data = doc.render_docx(_com001())
    assert isinstance(data, bytes) and len(data) > 0


def test_render_docx_round_trips_and_has_content():
    reopened = _reopen(doc.render_docx(_com001()))
    assert reopened.paragraphs[0].text.startswith("COM-001")
    headings = [p.text for p in reopened.paragraphs if p.style.name.startswith("Heading")]
    for marker in ("Business", "KPIs & Meaning", "Technical", "Row-Level Security"):
        assert marker in headings, f"missing DOCX section: {marker}"


def test_render_docx_lists_every_measure():
    m = _com001()
    reopened = _reopen(doc.render_docx(m))
    cell_text = " ".join(c.text for t in reopened.tables for r in t.rows for c in r.cells)
    for table in m.semantic.tables:
        for measure in table.measures:
            assert measure.name in cell_text


def test_render_docx_deterministic_content():
    """Content/styling deterministic (I2); bytes are not (OOXML timestamps) —
    assert on extracted text, not the raw bytes (see docx_document.py)."""
    m = _com001()
    text1 = [p.text for p in _reopen(doc.render_docx(m)).paragraphs]
    text2 = [p.text for p in _reopen(doc.render_docx(m)).paragraphs]
    assert text1 == text2


def test_render_docx_raises_docx_render_error_on_bad_brand_spec(tmp_path):
    bad_spec = tmp_path / "bad.yaml"
    bad_spec.write_text("identity:\n  brand_id: x\n", encoding="utf-8")  # missing required fields
    with pytest.raises(doc.DocxRenderError):
        doc.render_docx(_com001(), brand_spec_path=bad_spec)


def test_cli_writes_docx_file(tmp_path: Path):
    out = tmp_path / "doc.docx"
    assert doc.main([str(COM001), "--format", "docx", "--out", str(out)]) == 0
    assert out.exists()
    reopened = Document(out)
    assert reopened.paragraphs[0].text.startswith("COM-001")


def test_cli_docx_falls_back_to_markdown_on_broken_brand_spec(tmp_path: Path):
    bad_spec = tmp_path / "bad.yaml"
    bad_spec.write_text("identity:\n  brand_id: x\n", encoding="utf-8")
    out = tmp_path / "doc.docx"
    assert doc.main([str(COM001), "--format", "docx", "--brand-spec", str(bad_spec), "--out", str(out)]) == 0
    # DoD: broken branding falls back to Markdown, never a crash or a silent empty file.
    assert not out.exists()
    fallback = out.with_suffix(".md")
    assert fallback.exists() and fallback.read_text(encoding="utf-8").startswith("# COM-001")


def test_cli_docx_requires_out():
    assert doc.main([str(COM001), "--format", "docx"]) == 1
